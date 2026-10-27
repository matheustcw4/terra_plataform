"""Fontes sintéticas para os testes de adapter: uma planilha no layout do template e uma API Metryx simulada.

A mesma fazenda (docs/01 §5.1) é descrita nas duas fontes para o teste de equivalência:
BEZERRO com 1.000 cab no fim de jun/2025; em julho +120 nascimentos, −5 mortes, −300 desmame (-> DESMAMA MACHO).
"""

import json
from datetime import datetime
from pathlib import Path

import httpx
from openpyxl import Workbook

from app.etl.adapters.metryx.endpoints import ENDPOINTS

CABECALHOS = {
    "fUnidade": ["PROJETO", "PROPRIETÁRIO", "FAZENDA", "CÓDIGO PROJETO", "SISTEMA PRODUTIVO", "RETIRO", "ANO"],
    "fCompras": ["OPERAÇÃO", "DATA", "ORIGEM", "CATEG_FAZENDA", "QTD", "PESO_VIVO", "TOTAL_PESO", "VALOR_CAB",
                 "VALOR TOTAL", "VALO_KG", "VALOR_ARROBA", "DATA_ABREV", "CATEGORIA_BI"],
    "fVendaemPe": ["OPERAÇÃO", "CLIENTE", "DATA", "DOCUMENTO", "CATEG_FAZENDA", "QTD", "PESO_VIVO", "PESOVIVO_TOTAL",
                   "VALOR_ARROBA", "VALOR_CAB", "FATURAMENTO"],
    "fAbate": ["OPERAÇÃO", "CLIENTE", "CATEG_FAZENDA", "DATA", "QTD", "ERA", "PESO_VIVO", "RENDIMENTO", "PESO_ARROBA",
               "VALOR_ARROBA", "FATURAMENTO", "PESOVIVO_TOTAL", "PESOMORTO_TOTAL"],
    "fMortes": ["OPERAÇÃO", "DATA", "CATEG_FAZENDA", "CARIMBO_BEZ", "QTD", "CAUSA", "LOCAL", "PESO_VIVO", "SAFRA"],
    "fNascimento": ["OPERAÇÃO", "DATA", "CATEG_FAZENDA", "QTD", "SAFRA", "PESO_NASC"],
    "fdesmame": ["OPERAÇÃO", "DATA", "CATEGORIA_FAZENDA_SAÍDA", "CATEGORIA_FAZENDA_ENTRADA", "QTD", "PESO_VIVO"],
    "fRebanhoMedio": ["GRUPO_CATEGORIA", "QTD_CAB", "QTD_UA", "DATA"],
    "fArea": ["DATA", "ÁREA_PASTO", "ÁREA_ILP", "ÁREA_TOTAL"],
    "fDre": ["CENTRO_DE_CUSTO", "SUB_CENTRO", "VALOR", "DATA_PAG", "OPERAÇÃO", "CENTRO_DE_CUSTO_BI", "SUB_CENTRO_BI"],
}
CABECALHOS["fAbatePLANEJADO"] = CABECALHOS["fAbate"]          # metas: mesma aba com o sufixo PLANEJADO
CABECALHOS["fComprasPLANEJADO"] = CABECALHOS["fCompras"]
CABECALHOS["fNascimentoPLANEJADO"] = CABECALHOS["fNascimento"]
CABECALHOS["fMortesPLANEJADO"] = CABECALHOS["fMortes"]
CABECALHOS["ftranfentrada"] = ["OPERAÇÃO", "DATA", "CATEG_FAZENDA", "QTD", "PESO_VIVO", "TOTAL_PESO", "VALOR TOTAL"]
CABECALHOS["ftranssaida"] = CABECALHOS["ftranfentrada"]
CABECALHOS["fDrePLANEJADO"] = CABECALHOS["fDre"]
CABECALHOS["fRebanhoMedioPLANEJADO"] = CABECALHOS["fRebanhoMedio"]
PERIODOS = [("SAFRA 24/25", datetime(2024, 7, 1), datetime(2024, 6, 1)),
            ("SAFRA 25/26", datetime(2025, 7, 1), datetime(2025, 6, 1)),
            ("SAFRA 26/27", datetime(2026, 7, 1), datetime(2026, 6, 1)),
            ("CIVIL 2025", datetime(2025, 1, 1), datetime(2024, 12, 1)),
            ("CIVIL 2026", datetime(2026, 1, 1), datetime(2025, 12, 1)),
            ("CIVIL 2027", datetime(2027, 1, 1), datetime(2026, 12, 1))]


def gerar_planilha(pasta: Path, nome: str = "RELATÓRIO - 85 - TESTE.xlsx", *, codigo: int = 85,
                   fazenda: str = "TESTE", periodo: str = "SAFRA 25/26", saldo_bezerro: int = 1000,
                   linhas: dict[str, list[list]] | None = None,
                   categorias_extras: tuple[tuple[str, str, int | None], ...] = ()) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for aba in ["ConsultaPlanoContas", "dCategorias", "BASEPLANODECONTAS", "fUnidade", "CADASTROINICIAL",
                *[a for a in CABECALHOS if a != "fUnidade"]]:
        wb.create_sheet(aba)
    wb["fUnidade"].append(CABECALHOS["fUnidade"])
    wb["fUnidade"].append(["PROJETO TESTE", "FULANO DE TAL", fazenda, codigo, "CICLO COMPLETO", fazenda, "SAFRA"])

    ci = wb["CADASTROINICIAL"]
    ci["A2"], ci["B2"] = "PROJETO", "PROJETO TESTE"
    ci["A8"], ci["B8"] = "PERÍODO", periodo
    for i, (rotulo, inicio, anterior) in enumerate(PERIODOS):
        col = 15 + i
        ci.cell(9, col, rotulo)
        ci.cell(10, col, inicio)
        ci.cell(22, col, anterior)
    for c, valor in zip("ABCDEF", ("CATEGORIA FAZENDA", "CATEGORIA BI", "SEXO", "GRUPO CATEGORIA", "ERA",
                                   "SALDO INICIAL"), strict=True):
        ci[f"{c}10"] = valor
    cadastro = [("Bezerro Mamando", "BEZERRO", saldo_bezerro), ("Desmama Macho", "MACHOS 8-12", None),
                ("Vaca", "MATRIZES", 500), *categorias_extras]
    for n, (local, bi, saldo) in enumerate(cadastro, start=11):
        ci[f"A{n}"], ci[f"B{n}"], ci[f"F{n}"] = local, bi, saldo

    bp = wb["BASEPLANODECONTAS"]
    bp.append(["CENTRO DE CUSTO - INTTEGRA", "CENTRO DE CUSTO - METRYX", None, "SUBCENTRO - INTTEGRA",
               "SUBCENTRO - METRYX", "SUBCENTRO - METRYX"])
    bp.append(["PASTO", "Pastagens", None, "ADUBO", "Adubação", "Adubação"])
    cp = wb["ConsultaPlanoContas"]
    cp.append(["Plano de Contas - Débitos"])
    cp.append(["4.1.1 - Adubação", "Adubação", "DCM", "CCM", "CV", "D"])

    padrao = {
        "fNascimento": [["NASCIMENTO", datetime(2025, 7, 10), "Bezerro Mamando", 120, "25/26", 32]],
        "fMortes": [["MORTE", datetime(2025, 7, 12), "Bezerro Mamando", None, 5, "Cobra", "Pasto 1", 40, "25/26"]],
        "fdesmame": [["DESMAME", datetime(2025, 7, 20), "Bezerro Mamando", "Desmama Macho", 300, 200]],
        "fVendaemPe": [["EMPÉ", None, datetime(2025, 8, 5), "GTA1", "Vaca", 10, 450, 4500, None, None, 30000]],
        "fRebanhoMedio": [["Bezerros", 1050, None, datetime(2025, 7, 1)], ["Adultos", 500, None, datetime(2025, 7, 1)]],
        "fArea": [[datetime(2025, 7, 1), 800, 200, 1000]],
        "fDre": [["PASTO", "ADUBO", 1500.5, datetime(2025, 7, 15), "D", "Pastagens", "Adubação"],
                 ["XPTO LOCAL", "QUALQUER", 99, datetime(2025, 7, 16), None, None, None],
                 ["Apicultura", "Mel", 200, datetime(2025, 7, 17), "C", "Apicultura", "Mel"]],
    }
    for aba, dados in (padrao | (linhas or {})).items():
        ws = wb[aba]
        ws.append(CABECALHOS[aba])
        for linha in dados:
            ws.append(linha)
    caminho = pasta / nome
    wb.save(caminho)
    return caminho


# --- API Metryx simulada -------------------------------------------------------------------------------------------

CLIENTE = "c0ffee00-0000-0000-0000-000000000001"
FAZENDA = "fa000000-0000-0000-0000-000000000001"
FAZENDA_2 = "fa000000-0000-0000-0000-000000000002"
RETIRO = "re000000-0000-0000-0000-000000000001"
CAT_BEZ, CAT_DESM, CAT_VACA = "ca000001", "ca000002", "ca000003"


def pastagem(id_pastagem: str, utilizacao: str, area: float, inicio: str, fim: str | None = None,
             fazenda: str = FAZENDA, produtiva: float | None = None) -> dict:
    """Um período de uso de uma pastagem, como o `utilizacao-pasto` devolve (sem idCustomer)."""
    if produtiva is None:
        produtiva = area
    return {"fazendaId": fazenda, "fazendaNome": "Fazenda Metryx", "retiroId": 77, "retiroName": "Sede",
            "moduloId": "m1", "moduloNome": "Módulo 1", "id": id_pastagem, "name": f"Pasto {id_pastagem}",
            "tipoUsoId": f"{id_pastagem}-{inicio}", "utilizacaoArea": utilizacao, "statusPecuaria": "Disponível",
            "dataInicio": inicio, "dataFim": fim, "areaCalculada": area, "areaProdutiva": produtiva,
            "areaTotal": area}


def payloads_metryx(**sobrescrever) -> dict[str, list[dict]]:
    """Mesma fazenda do doc 01 §5.1 no formato da API (campos do OpenAPI)."""
    base = {"idCustomer": CLIENTE, "customer": "Cliente Teste", "idFarm": FAZENDA, "farmName": "Fazenda Metryx",
            "idRetiro": RETIRO, "retiroName": "Sede", "owner": "Dono"}
    bal = lambda d, cat, q, **kw: {**base, "date": d, "uuidCategory": cat, "amount": q, **kw}   # noqa: E731
    dados = {
        "fazenda": [{"idCustomer": CLIENTE, "customer": "Cliente Teste", "idFarm": FAZENDA, "farmName": "Fazenda Metryx"},
                    {"idCustomer": CLIENTE, "customer": "Cliente Teste", "idFarm": FAZENDA_2, "farmName": "Outra"}],
        "categoria": [
            {"idCustomer": CLIENTE, "idCategory": CAT_BEZ, "category": "Bezerro", "categoryGroup": "Bezerros Mamando",
             "ageCategory": "00-08", "gender": "Macho"},
            {"idCustomer": CLIENTE, "idCategory": CAT_DESM, "category": "Garrote 8-12", "categoryGroup":
             "Garrotes e Bois", "ageCategory": "08-12", "gender": "Macho"},
            {"idCustomer": CLIENTE, "idCategory": CAT_VACA, "category": "Vaca", "categoryGroup":
             "Matrizes em Reprodução", "ageCategory": "36-00", "gender": "Fêmea"}],
        "retiro": [{**base}],
        # saldo de fim de mês = número oficial do Metryx; julho fecha com 815 (sem ajuste)
        "balanco": [bal("2025-06-30", CAT_BEZ, 1000, birth="06/25", harvest="24/25"),
                    bal("2025-07-31", CAT_BEZ, 815, birth="06/25", harvest="24/25"),
                    bal("2025-06-30", CAT_DESM, 0), bal("2025-07-31", CAT_DESM, 300),
                    bal("2025-06-30", CAT_VACA, 500), bal("2025-08-31", CAT_VACA, 490)],
        "nascimento": [{**base, "date": "2025-07-10", "idCategory": CAT_BEZ, "amount": 120, "harvest": "25/26"}],
        "morte": [{**base, "idRetiro": "77", "date": "2025-07-12", "categoryId": CAT_BEZ, "amount": 5,
                   "cause": "Cobra", "totalWeight": 200}],                      # idRetiro legado (inteiro)
        "desmame": [{**base, "date": "2025-07-20", "oldCategoryId": CAT_BEZ, "newCategoryId": CAT_DESM, "amount": 300,
                     "weight": 200}],
        "venda": [{**base, "date": "2025-08-05", "categoryId": CAT_VACA, "amount": 10, "weightTotal": 4500,
                   "totalFaturamento": 30000, "client": "Comprador", "deadArrobaAvg": None}],
        "compra": [],
        "mudanca_categoria": [],
        "transferencia": [{"id": "t1", "itemId": "i1", "idCustomer": CLIENTE, "idFarmOrigin": FAZENDA,
                           "idRetiroOrigin": RETIRO, "retiroNameOrigin": "Sede", "idFarmDestiny": FAZENDA_2,
                           "farmNameDestiny": "Outra", "date": "2025-09-01", "categoryId": CAT_VACA, "amount": 3,
                           "weightTotal": 1350, "totalFaturamento": 0}],
        "rebanho_medio": [{**base, "month": "7-2025", "idCategory": CAT_BEZ, "amount": 1050},
                          {**base, "month": "7-2025", "idCategory": CAT_VACA, "amount": 500}],
        # a área da planilha sintética (jul/2025: pasto 800 + ILP 200 = 1.000 ha) em duas pastagens de uso Pecuária;
        # a APP não entra na área de pasto
        "pasto": [pastagem("p1", "Pecuária", 800, "2025-01-01"), pastagem("p2", "Pecuária", 200, "2025-01-01"),
                  pastagem("p3", "APP", 50, "2025-01-01")],
        "lancamento": [
            {"idCustomer": CLIENTE, "idFarmFiscal": FAZENDA, "idfarmDestination": FAZENDA, "idFinancialEntrie": "f1",
             "idPayment": "p1", "parcel": "1", "costCenter": "Pastagens​", "subCenter": "Adubação",
             "type": "Débito", "amountPaid": 1500.5, "datePayment": "2025-07-15", "dueDate": "2025-07-10"},
            {"idCustomer": CLIENTE, "idFarmFiscal": FAZENDA, "idfarmDestination": FAZENDA, "idFinancialEntrie": "f2",
             "idPayment": None, "parcel": "1", "costCenter": "Pastagens", "subCenter": "Adubação",
             "type": "Débito", "amountPaid": 0, "datePayment": None, "dueDate": "2025-08-10"}],   # não pago
    }
    dados.update(sobrescrever)
    return dados


def transporte_metryx(dados: dict[str, list[dict]], status: dict[str, int] | None = None,
                      chamadas: list[str] | None = None, token: str = "token-teste") -> httpx.MockTransport:
    caminho_para_nome = {e.caminho: e.nome for e in ENDPOINTS.values()}

    def responder(req: httpx.Request) -> httpx.Response:
        if chamadas is not None:
            chamadas.append(req.url.path)
        if req.headers.get("x-auth-token") != token:
            return httpx.Response(401)
        nome = caminho_para_nome.get(req.url.path.removeprefix("/mtxserver"))
        if nome is None:
            return httpx.Response(404)
        if status and nome in status:
            return httpx.Response(status[nome])
        return httpx.Response(200, content=json.dumps(dados.get(nome, [])).encode())

    return httpx.MockTransport(responder)
