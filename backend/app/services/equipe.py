"""A equipe da consultoria a partir da planilha de levantamento dos clientes (pedido do usuário, 15/10/2026).

Cada linha da planilha é uma fazenda, com o técnico (nome e e-mail), o e-mail do gerente e o do gerente regional.
Daqui saem:
- os usuários, identificados pelo e-mail. O nome vem da coluna TÉCNICO; quem só aparece como gerente ou regional
  ganha um nome provisório tirado do e-mail (o admin corrige na tela de usuários);
- a função: "Gerente" para quem aparece como gerente ou regional em alguma linha (o regional é gerente igual);
  "Técnico" para os demais;
- as fazendas que cada um vê: as das linhas em que aparece, em qualquer das três colunas.

Como a linha casa com as fazendas do painel:
- planilha: pelo código do cliente e o nome da fazenda, a mesma identidade do adapter ("código|FAZENDA"); se o código
  tiver uma fazenda só, vale ela, mesmo com outro nome (com aviso);
- Metryx: os nomes das fazendas no Metryx não seguem os da planilha ("3) Fazenda Itauba" × "ITAÚBA"), então o
  vínculo é pelo cliente: proprietário da linha igual ao customer do Metryx dá acesso às fazendas Metryx dele (no
  levantamento, cada cliente tem uma equipe só).

A gravação só acrescenta: não apaga usuário nem acesso e não troca o nome, a senha nem a função que o admin já
escolheu. Usuário novo nasce sem senha e não entra até o admin definir uma na tela de usuários.
"""

import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import Connection, text

from app.etl.adapters.planilha.template import norm, norm_texto

TECNICO = "Técnico"
GERENTE = "Gerente"
# campo -> a coluna da planilha, normalizada (sem acento, maiúsculas, "_" no lugar de espaço e hífen)
COLUNAS = {"codigo": "CODIGO_DO_CLIENTE", "proprietario": "PROPRIETARIO", "fazenda": "FAZENDA", "tecnico": "TECNICO",
           "email_tecnico": "EMAIL", "email_gerente": "EMAIL_GERENTE", "email_regional": "EMAIL_REGIONAL"}
PARTICULAS = {"da", "de", "do", "das", "dos", "e"}


@dataclass(frozen=True)
class LinhaDoLevantamento:
    numero: int                             # a linha na planilha, para os avisos
    codigo: str
    proprietario: str
    fazenda: str
    tecnico: str
    email_tecnico: str
    email_gerente: str
    email_regional: str


@dataclass(frozen=True)
class FazendaDoPainel:
    id: UUID
    fonte: str                              # planilha | metryx
    id_externo: str
    proprietario: str


@dataclass
class Pessoa:
    email: str
    nome: str = ""
    nome_provisorio: bool = False           # tirado do e-mail: o admin confere na tela de usuários
    funcao: str = TECNICO
    fazendas: set[UUID] = field(default_factory=set)


@dataclass
class Equipe:
    pessoas: dict[str, Pessoa] = field(default_factory=dict)
    avisos: list[str] = field(default_factory=list)
    linhas_sem_fazenda: list[LinhaDoLevantamento] = field(default_factory=list)
    clientes_metryx_fora: list[str] = field(default_factory=list)   # customers do painel sem linha no levantamento


# --- leitura da planilha -------------------------------------------------------------------------------------------

def ler_levantamento(caminho: Path) -> list[LinhaDoLevantamento]:
    livro = load_workbook(caminho, read_only=True, data_only=True)
    linhas = livro.worksheets[0].iter_rows(values_only=True)
    posicao = posicoes_das_colunas([norm(valor) for valor in next(linhas)])
    resultado = []
    for numero, valores in enumerate(linhas, start=2):
        if any(valores):
            resultado.append(linha_do_levantamento(numero, valores, posicao))
    livro.close()
    return resultado


def posicoes_das_colunas(cabecalho: list[str]) -> dict[str, int]:
    posicao = {}
    for campo, coluna in COLUNAS.items():
        if coluna not in cabecalho:
            raise ValueError(f"a planilha não tem a coluna {coluna}")
        posicao[campo] = cabecalho.index(coluna)
    return posicao


def linha_do_levantamento(numero: int, valores: tuple, posicao: dict[str, int]) -> LinhaDoLevantamento:
    campos = {}
    for campo, indice in posicao.items():
        valor = ""
        if indice < len(valores):
            valor = para_texto(valores[indice])
        if campo.startswith("email"):
            valor = valor.lower()
        campos[campo] = valor
    return LinhaDoLevantamento(numero=numero, **campos)


def para_texto(valor) -> str:
    """O código vem como número do Excel (11 ou 11.0): vira "11"."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


# --- nomes ---------------------------------------------------------------------------------------------------------

def nome_proprio(texto: str) -> str:
    """"CARLOS CHULLI" -> "Carlos Chulli"; "MARIA DA SILVA" -> "Maria da Silva"."""
    palavras = []
    for palavra in texto.split():
        if palavra.lower() in PARTICULAS and palavras:
            palavras.append(palavra.lower())
        else:
            palavras.append(palavra.capitalize())
    return " ".join(palavras)


def nome_do_email(email: str) -> str:
    """Nome provisório: "joao.silva@x" -> "Joao Silva"; "lucasbraido@x" -> "Lucasbraido"."""
    partes = re.split(r"[._\-]+", email.split("@")[0])
    return " ".join(parte.capitalize() for parte in partes if parte)


def peso_do_nome(item: tuple[str, int]) -> tuple[int, int]:
    """O nome mais usado; no empate, o mais completo."""
    nome, quantas = item
    return quantas, len(nome)


# --- montagem ------------------------------------------------------------------------------------------------------

def montar_equipe(linhas: list[LinhaDoLevantamento], fazendas: list[FazendaDoPainel]) -> Equipe:
    montagem = MontagemDaEquipe(linhas, fazendas)
    montagem.pessoas_e_funcoes()
    montagem.nomes()
    montagem.fazendas()
    montagem.clientes_metryx_sem_linha()
    return montagem.equipe


class MontagemDaEquipe:
    def __init__(self, linhas: list[LinhaDoLevantamento], fazendas: list[FazendaDoPainel]):
        self.linhas = linhas
        self.equipe = Equipe()
        self.planilha_por_chave: dict[str, UUID] = {}
        self.planilha_por_codigo: dict[str, list[UUID]] = defaultdict(list)
        self.metryx_por_cliente: dict[str, list[UUID]] = defaultdict(list)
        for fazenda in fazendas:
            if fazenda.fonte == "planilha":
                self.planilha_por_chave[fazenda.id_externo] = fazenda.id
                self.planilha_por_codigo[fazenda.id_externo.split("|")[0]].append(fazenda.id)
            elif fazenda.fonte == "metryx":
                self.metryx_por_cliente[norm_texto(fazenda.proprietario)].append(fazenda.id)

    def pessoa(self, email: str) -> Pessoa:
        if email not in self.equipe.pessoas:
            self.equipe.pessoas[email] = Pessoa(email=email)
        return self.equipe.pessoas[email]

    def pessoas_e_funcoes(self) -> None:
        """Gerente ou regional em alguma linha é "Gerente", mesmo que seja técnico em outras."""
        for linha in self.linhas:
            if linha.email_tecnico:
                self.pessoa(linha.email_tecnico)
            for email in (linha.email_gerente, linha.email_regional):
                if email:
                    self.pessoa(email).funcao = GERENTE

    def nomes(self) -> None:
        """O nome mais usado na coluna TÉCNICO para o e-mail. Se ele for o nome de outra pessoa (que o usa em mais
        linhas), a linha está trocada: o nome sai do e-mail. Quem não é técnico em linha nenhuma também."""
        nomes_por_email: dict[str, Counter] = defaultdict(Counter)
        for linha in self.linhas:
            if linha.email_tecnico and linha.tecnico:
                nomes_por_email[linha.email_tecnico][nome_proprio(linha.tecnico)] += 1
        escolhidos = self.nomes_escolhidos(nomes_por_email)
        for email, pessoa in self.equipe.pessoas.items():
            if email in escolhidos:
                pessoa.nome = escolhidos[email]
            else:
                pessoa.nome = nome_do_email(email)
                pessoa.nome_provisorio = True
        self.avisar_linhas_trocadas(escolhidos)

    def nomes_escolhidos(self, nomes_por_email: dict[str, Counter]) -> dict[str, str]:
        """Primeiro tira os nomes que são de outra pessoa (quem o usa em mais linhas); depois fica o mais usado."""
        dono_do_nome: dict[str, tuple[int, str]] = {}           # nome -> (linhas, e-mail de quem mais o usa)
        for email, nomes in nomes_por_email.items():
            for nome, quantas in nomes.items():
                if nome not in dono_do_nome or quantas > dono_do_nome[nome][0]:
                    dono_do_nome[nome] = (quantas, email)
        escolhidos = {}
        for email, nomes in nomes_por_email.items():
            proprios = {}
            for nome, quantas in nomes.items():
                if dono_do_nome[nome][1] == email:
                    proprios[nome] = quantas
            if proprios:
                escolhidos[email] = max(proprios.items(), key=peso_do_nome)[0]
            else:
                self.equipe.avisos.append(f"{email}: o nome escrito ({', '.join(nomes)}) é de outra pessoa; ficou o "
                                          "nome tirado do e-mail")
        return escolhidos

    def avisar_linhas_trocadas(self, escolhidos: dict[str, str]) -> None:
        """Linha com o nome de outra pessoa no e-mail de alguém. A forma curta do mesmo nome ("Carlos" para
        "Carlos Chulli") não é troca."""
        for linha in self.linhas:
            nome = nome_proprio(linha.tecnico)
            escolhido = escolhidos.get(linha.email_tecnico)
            if escolhido is None or not nome or escolhido.startswith(nome):
                continue
            self.equipe.avisos.append(f"linha {linha.numero} ({linha.fazenda}): o técnico escrito é '{nome}', mas "
                                      f"o e-mail é de {escolhido} ({linha.email_tecnico}); valeu o e-mail")

    def fazendas(self) -> None:
        for linha in self.linhas:
            ids = self.fazendas_da_linha(linha)
            if not ids:
                self.equipe.linhas_sem_fazenda.append(linha)
                continue
            for email in (linha.email_tecnico, linha.email_gerente, linha.email_regional):
                if email:
                    self.pessoa(email).fazendas.update(ids)

    def fazendas_da_linha(self, linha: LinhaDoLevantamento) -> set[UUID]:
        ids = set(self.metryx_por_cliente.get(norm_texto(linha.proprietario), []))
        chave = f"{linha.codigo}|{norm_texto(linha.fazenda)}"
        do_codigo = self.planilha_por_codigo.get(linha.codigo, [])
        if chave in self.planilha_por_chave:
            ids.add(self.planilha_por_chave[chave])
        elif len(do_codigo) == 1:
            ids.add(do_codigo[0])
            self.equipe.avisos.append(f"linha {linha.numero}: '{linha.fazenda}' (código {linha.codigo}) casada pela "
                                      "única fazenda de planilha do código, que tem outro nome")
        return ids

    def clientes_metryx_sem_linha(self) -> None:
        clientes_do_levantamento = {norm_texto(linha.proprietario) for linha in self.linhas}
        for cliente in sorted(self.metryx_por_cliente):
            if cliente not in clientes_do_levantamento:
                self.equipe.clientes_metryx_fora.append(cliente)


# --- leitura do painel e gravação ----------------------------------------------------------------------------------

FAZENDAS_DO_PAINEL = """
    SELECT f.id, v.fonte, v.id_externo, coalesce(f.proprietario, '') AS proprietario
    FROM core.fonte_vinculo v JOIN core.fazenda f ON f.id = v.entidade_id
    WHERE v.entidade = 'fazenda'
"""
FUNCAO = "SELECT id FROM core.funcao WHERE lower(nome) = lower(:nome)"
CRIAR_FUNCAO = "INSERT INTO core.funcao (nome) VALUES (:nome) RETURNING id"
USUARIO = "SELECT id, funcao_id FROM core.usuario WHERE email = lower(btrim(:email))"
CRIAR_USUARIO = """
    INSERT INTO core.usuario (email, nome, funcao_id) VALUES (lower(btrim(:email)), :nome, :funcao) RETURNING id
"""
COMPLETAR_FUNCAO = "UPDATE core.usuario SET funcao_id = :funcao, atualizado_em = now() WHERE id = :id"
DAR_ACESSOS = """
    INSERT INTO core.usuario_fazenda (usuario_id, fazenda_id)
    SELECT :usuario, fazenda FROM unnest(CAST(:fazendas AS uuid[])) AS fazenda
    ON CONFLICT DO NOTHING
"""


def fazendas_do_painel(conexao: Connection) -> list[FazendaDoPainel]:
    fazendas = []
    for linha in conexao.execute(text(FAZENDAS_DO_PAINEL)).mappings():
        fazendas.append(FazendaDoPainel(**linha))
    return fazendas


def aplicar_equipe(conexao: Connection, equipe: Equipe) -> Counter:
    """Grava a equipe e conta o que mudou: usuarios_criados, funcoes_completadas, acessos_novos."""
    funcoes = {TECNICO: garantir_funcao(conexao, TECNICO), GERENTE: garantir_funcao(conexao, GERENTE)}
    resumo = Counter()
    for email in sorted(equipe.pessoas):
        pessoa = equipe.pessoas[email]
        usuario_id = garantir_usuario(conexao, pessoa, funcoes[pessoa.funcao], resumo)
        if pessoa.fazendas:
            novos = conexao.execute(text(DAR_ACESSOS), {"usuario": usuario_id,
                                                        "fazendas": [str(fazenda) for fazenda in pessoa.fazendas]})
            resumo["acessos_novos"] += novos.rowcount
    return resumo


def garantir_funcao(conexao: Connection, nome: str) -> int:
    funcao_id = conexao.scalar(text(FUNCAO), {"nome": nome})
    if funcao_id is None:
        funcao_id = conexao.scalar(text(CRIAR_FUNCAO), {"nome": nome})
    return funcao_id


def garantir_usuario(conexao: Connection, pessoa: Pessoa, funcao_id: int, resumo: Counter) -> UUID:
    """Usuário que já existe fica como está (nome, senha, admin); só ganha a função se ainda não tiver uma."""
    existente = conexao.execute(text(USUARIO), {"email": pessoa.email}).first()
    if existente is None:
        resumo["usuarios_criados"] += 1
        return conexao.scalar(text(CRIAR_USUARIO), {"email": pessoa.email, "nome": pessoa.nome, "funcao": funcao_id})
    if existente.funcao_id is None:
        conexao.execute(text(COMPLETAR_FUNCAO), {"id": existente.id, "funcao": funcao_id})
        resumo["funcoes_completadas"] += 1
    return existente.id
